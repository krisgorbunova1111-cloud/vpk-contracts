"""Word-документы исполнительной документации: папка ИД (титул + реестр формы 1.1 +
АОСР + реестр документов о качестве), отдельный АОСР, отчёт по бланкам проверки.

Сканы (приказы, сертификаты, схемы, журналы) в Word не вставляются — на их месте
стоит лист-вкладыш с названием документа, архивариус подкладывает оригинал.
"""
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from . import id_docs
from .docgen import _center, _doc, _save, _table
from .id_docs import ACCESS, BLANKS_BY_ID, RESULT_KEYS, org, person, roles_of, ru_date
from .seed import AOSR_ROLES

MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
          "сентября", "октября", "ноября", "декабря"]


def _field(doc, value, caption, bold_label=None):
    """Строка «значение» с мелкой подписью под ней — как в бланке АОСР."""
    if bold_label:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(bold_label).bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(value or " ")
    r.italic, r.underline = True, True
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(4)
    cr = c.add_run(f"({caption})")
    cr.font.size = Pt(7)


def _org_line(g):
    if not g:
        return ""
    parts = [g["name"], f"ОГРН {g['ogrn']}", f"ИНН {g['inn']}", g["address"], g.get("phone", "")]
    if g.get("sro"):
        parts.append(g["sro"])
    return ", ".join(p for p in parts if p and p != "—")


def _person_line(db, pid):
    p = person(db, pid)
    if not p:
        return ""
    g = org(db, p["org_id"])
    s = f"{p['position']} {g['name'] if g else ''} {p['fio']}"
    if p.get("nrs"):
        s += f", идентификационный номер в НРС {p['nrs']}"
    if p.get("order_no"):
        s += f", приказ № {p['order_no']} от {ru_date(p.get('order_date'))}"
    return s


def _short_fio(db, pid):
    p = person(db, pid)
    if not p:
        return ""
    f = p["fio"].split()
    return f"{f[0]} {''.join(x[0] + '.' for x in f[1:])}"


def _date_words(iso):
    if not iso:
        return "«___» __________ 20__ г."
    y, m, d = iso.split("-")
    return f"«{d}» {MONTHS[int(m) - 1]} {y} г."


def add_aosr(doc, db, obj, folder, a):
    r = roles_of(db, obj["id"])
    _field(doc, f"{obj['name']}, {obj['address']}",
           "наименование объекта капитального строительства в соответствии с проектной документацией, почтовый или строительный адрес",
           "Объект капитального строительства")
    _field(doc, _org_line(org(db, r["orgs"].get("developer"))),
           "фамилия, имя, отчество, адрес места жительства, ОГРН, ИНН, наименование, номер телефона/факса",
           "Застройщик (технический заказчик, эксплуатирующая организация или региональный оператор)")
    _field(doc, _org_line(org(db, r["orgs"].get("builder"))),
           "наименование, ОГРН, ИНН, место нахождения, номер телефона/факса, наименование СРО",
           "Лицо, осуществляющее строительство")
    _field(doc, _org_line(org(db, r["orgs"].get("designer"))),
           "наименование, ОГРН, ИНН, место нахождения, номер телефона/факса, наименование СРО",
           "Лицо, осуществляющее подготовку проектной документации")

    _center(doc, "АКТ", 13, True, 0)
    _center(doc, "освидетельствования скрытых работ", 12, True, 2)
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = f"№ {a['no']}"
    t.rows[0].cells[1].text = _date_words(a["date"])
    t.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    caption = "должность, фамилия, инициалы, реквизиты документа о представительстве"
    for role in AOSR_ROLES:
        _field(doc, _person_line(db, r["people"].get(role["id"])), caption, role["name"])
    others = [x for x in r["people"].get("others", []) if x]
    if others:
        _field(doc, "; ".join(_person_line(db, x) for x in others), caption,
               "а также иные представители лиц, участвующих в освидетельствовании:")

    ex = org(db, r["orgs"].get("executor"))
    _field(doc, ex["name"] if ex else "", "наименование лица, выполнившего работы, подлежащие освидетельствованию",
           "произвели осмотр работ, выполненных")
    doc.add_paragraph("и составили настоящий акт о нижеследующем:")

    mats = [m for m in db["materials"] if m["id"] in a["materials"]]
    mat_text = (", ".join(m["name"] for m in mats) + ". Документы о качестве согласно реестру — приложение к акту.") if mats else "—"
    s = a.get("scheme") or {}
    scheme = f"Исполнительная схема № {s['no']} от {ru_date(s.get('date'))} на {s.get('sheets', 1)} листах" if s.get("no") else "—"
    items = [
        ("1. К освидетельствованию предъявлены следующие работы:", a["works"], "наименование скрытых работ"),
        ("2. Работы выполнены по проектной документации", a.get("basis") or id_docs.basis_text(folder),
         "номер, другие реквизиты чертежа, наименование проектной и/или рабочей документации"),
        ("3. При выполнении работ применены", mat_text,
         "наименование строительных материалов (изделий), реквизиты сертификатов и/или других документов, подтверждающих их качество"),
        ("4. Предъявлены документы, подтверждающие соответствие работ предъявляемым требованиям:", scheme,
         "исполнительные схемы и чертежи, результаты экспертиз, обследований, лабораторных и иных испытаний"),
    ]
    for label, value, cap in items:
        _field(doc, value, cap, label)
    doc.add_paragraph(f"5. Даты: начала работ {_date_words(a['start'])}\n             окончания работ {_date_words(a['end'])}")
    _field(doc, "\n".join(a.get("norms", [])),
           "наименования и структурные единицы технических регламентов, иных нормативных правовых актов, разделы проектной и/или рабочей документации",
           "6. Работы выполнены в соответствии с")
    _field(doc, a.get("next_works", ""), "наименование работ, конструкций, участков сетей инженерно-технического обеспечения",
           "7. Разрешается производство последующих работ")
    doc.add_paragraph(f"Дополнительные сведения: {a.get('extra_info') or 'нет'}")
    doc.add_paragraph(f"Акт составлен в {a.get('copies', 4)} экземплярах.")
    apps = []
    if s.get("no"):
        apps.append(f"1. Исполнительная схема № {s['no']} от {ru_date(s.get('date'))} на {s.get('sheets', 1)} листах;")
    if mats:
        apps.append(f"{len(apps) + 1}. Реестр документов, подтверждающих качество материалов, с сертификатами и паспортами.")
    doc.add_paragraph("Приложения:\n" + ("\n".join(apps) or "нет"))
    for role in AOSR_ROLES:
        pid = r["people"].get(role["id"])
        if not pid:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.add_run(role["name"]).font.size = Pt(9)
        doc.add_paragraph(f"{_short_fio(db, pid)}  ____________________ (подпись)")
    for pid in others:
        doc.add_paragraph(f"Представитель иных лиц: {_short_fio(db, pid)}  ____________________ (подпись)")


def _folder_cover(doc, db, obj, folder):
    r = roles_of(db, obj["id"])
    for key, label in (("developer", "Заказчик"), ("builder", "Подрядчик"), ("executor", "Исполнитель работ")):
        g = org(db, r["orgs"].get(key))
        if g:
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(g["name"])
    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    p.add_run("Адрес объекта: ").bold = True
    p.add_run(obj["address"])
    p = doc.add_paragraph()
    p.add_run("Вид работ: ").bold = True
    p.add_run(folder["title"])
    doc.add_paragraph("\n\n\n")
    _center(doc, "ИСПОЛНИТЕЛЬНАЯ ДОКУМЕНТАЦИЯ", 20, True, 60)
    _center(doc, f"Папка № {folder['no']}", 16, True)
    doc.add_page_break()


def _registry_page(doc, db, obj, folder, rows):
    r = roles_of(db, obj["id"])
    p = doc.add_paragraph("Приложение № 2\n(Реестр исполнительной документации)\nФорма № 1.1")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    names = [org(db, r["orgs"].get(k)) for k in ("developer", "builder", "executor")]
    doc.add_paragraph("\n".join(g["name"] for g in names if g))
    q = doc.add_paragraph(f"Объект: {obj['name']}, расположенный по адресу: {obj['address']}")
    q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    aosrs = id_docs.folder_aosr(db, folder["id"])
    reg_date = max((a["date"] for a in aosrs), default=None)
    _center(doc, f"РЕЕСТР № {folder['no']} от {ru_date(reg_date)}", 12, True, 0)
    _center(doc, "исполнительной документации", 12, True, 0)
    _center(doc, f"по объекту {obj['name']}", 12, True)
    _table(doc, ["№ п/п", "Наименование документа", "№ чертежа, акта, разрешения, журнала и др.", "Кол-во листов", "Лист по списку"],
           [(x["n"], x["name"], x["no"], x["sheets"], x["pages"]) for x in rows])


def _insert_sheet(doc, row):
    """Лист-вкладыш: сюда архивариус подкладывает оригинал/скан документа."""
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("\n\n\n")
    _center(doc, f"[ Лист {row['pages']} ]", 11)
    _center(doc, row["name"], 14, True)
    _center(doc, f"№ {row['no']} · листов: {row['sheets']}", 11)
    _center(doc, "Вложить оригинал или заверенную копию документа", 10)


def build_folder(db, obj, folder) -> bytes:
    rows = id_docs.registry(db, folder)
    doc = _doc()
    _folder_cover(doc, db, obj, folder)
    _registry_page(doc, db, obj, folder, rows)
    aosr_by_id = {a["id"]: a for a in id_docs.folder_aosr(db, folder["id"])}
    for row in rows[1:]:
        if row.get("kind") == "aosr":
            doc.add_page_break()
            add_aosr(doc, db, obj, folder, aosr_by_id[row["id"]])
        elif row.get("kind") == "quality":
            doc.add_page_break()
            _center(doc, "РЕЕСТР документов, подтверждающих качество", 12, True)
            mats = id_docs.folder_materials(db, folder["id"])
            _table(doc, ["№", "Материал (изделие)", "Документ о качестве", "Срок действия", "Паспорт"],
                   [(i + 1, m["name"], m["cert_no"], f"до {ru_date(m['cert_until'])}", "есть" if m.get("passport") else "—")
                    for i, m in enumerate(mats)])
        else:
            _insert_sheet(doc, row)
    return _save(doc)


def build_aosr(db, obj, folder, a) -> bytes:
    doc = _doc()
    add_aosr(doc, db, obj, folder, a)
    return _save(doc)


def build_inspections(db, obj, items) -> bytes:
    doc = _doc()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    _center(doc, "Результаты проверки состава работ и контрольных замеров", 15, True, 2)
    _center(doc, f"{obj['name']}, {obj['address']}", 12, False, 14)
    _table(doc, ["Бланк", "Вид работ", "Дата", "Место", "Состав", "Объемы", "Характ.", "Скрытые"], [
        (BLANKS_BY_ID[i["blank_id"]]["code"], BLANKS_BY_ID[i["blank_id"]]["name"], ru_date(i["date"]), _place(i),
         *[RESULT_KEYS[k][1].get(i["results"].get(k), "—") for k in RESULT_KEYS]) for i in items])
    for i in items:
        b = BLANKS_BY_ID[i["blank_id"]]
        doc.add_page_break()
        doc.add_heading(f"{b['code']} {b['name']}", level=1)
        doc.add_paragraph(f"Раздел: {b['section']}.  Дата: {ru_date(i['date'])}.  Проверяющий: {i['inspector'] or '—'}.  "
                          f"Представитель подрядчика: {i['contractor_rep'] or '—'}")
        doc.add_paragraph(f"Место: {_place(i)}.  Доступность: {ACCESS.get(i['access'], '—')}.")
        if i.get("aosr_id"):
            a = next((x for x in db["aosr"] if x["id"] == i["aosr_id"]), None)
            if a:
                doc.add_paragraph(f"Сверка с АОСР № {a['no']} от {ru_date(a['date'])}")
        _table(doc, ["Показатель", "Ед.", "По документам", "Замер 1", "Замер 2", "Итог"],
               [(r["name"], r["unit"], r["doc"], r["m1"], r["m2"], r["total"]) for r in i["rows"]])
        for label, val in (("Расчёт объема и вычетов", i["calc"]), ("Примечания и выявленные расхождения", i["notes"])):
            if val:
                doc.add_paragraph().add_run(label + ": ").bold = True
                doc.add_paragraph(val)
        doc.add_paragraph("; ".join(f"{name}: {opts.get(i['results'].get(k), '—')}" for k, (name, opts) in RESULT_KEYS.items()))
        if i["action"]:
            doc.add_paragraph(f"Требуемое действие: {i['action']}. Ответственный: {i['responsible'] or '—'}. Срок: {ru_date(i['deadline']) if i['deadline'] else '—'}")
        doc.add_paragraph("\nПроверяющий ____________________     Представитель подрядчика ____________________")
    remarks = [i for i in items if id_docs.has_discrepancy(i)]
    doc.add_page_break()
    doc.add_heading("Журнал замечаний и повторной проверки", level=1)
    if remarks:
        _table(doc, ["Бланк и пункт", "Место и расхождение", "Действие и ответственный", "Срок и повторная проверка"], [
            (BLANKS_BY_ID[i["blank_id"]]["code"], f"{_place(i)}. {i['notes']}", f"{i['action']} / {i['responsible']}",
             f"{ru_date(i['deadline']) if i['deadline'] else '—'} / {'закрыто' if i['status'] == 'closed' else 'открыто'}")
            for i in remarks])
    else:
        doc.add_paragraph("Расхождений не выявлено.")
    return _save(doc)


def _place(i):
    p = i["place"]
    parts = [f"подъезд {p['entrance']}" if p["entrance"] else "", f"этаж {p['floor']}" if p["floor"] else "",
             f"кв. {p['apt']}" if p["apt"] else "", p["axes"]]
    return ", ".join(x for x in parts if x) or "—"
