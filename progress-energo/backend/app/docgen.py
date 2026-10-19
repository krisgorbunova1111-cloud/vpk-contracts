"""Генерация документов Word: том ИД, заготовка раздела РД, заявка на допуск.

В прототипе документы собираются «с нуля» кодом. В боевой версии сюда
подключаются утверждённые шаблоны .docx с полями ({{адрес}}, {{дата}} …).
"""
import io
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (пригодится для листов А3)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

from .seed import BUILDING_TYPES, SYSTEMS

COMPANY = "ООО ВПК «Прогресс Энерго»"
STATE_RU = {"satisfactory": "Работоспособное", "limited": "Ограниченно работоспособное",
            "emergency": "Аварийное", "": "Не обследовано"}
CONSTR_RU = {"foundation": "Фундаменты", "walls": "Стены", "roof": "Кровля / чердак", "slabs": "Перекрытия"}

# Порядок комплектования тома ИД (по РД-11-02-2006, упрощённо)
ID_ORDER = [
    "Титульный лист",
    "Опись документов тома",
    "Акт технического обследования объекта",
    "Ведомость инженерных систем и узлов",
    "Обмерные планы (по данным лазерного сканирования)",
    "Ведомость применённых материалов",
    "Сертификаты и паспорта на материалы",
    "Протоколы лабораторных испытаний",
]


def _doc() -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for s in doc.sections:
        s.left_margin, s.right_margin = Cm(2.0), Cm(1.0)
        s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5)
    return doc


def _center(doc, text, size=12, bold=False, space=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    return p


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()
    return t


def _title_page(doc, obj, volume_title, volume_code):
    _center(doc, COMPANY, 14, True, 2)
    _center(doc, "Исполнитель работ", 10, False, 60)
    _center(doc, obj["customer"], 12, False, 2)
    _center(doc, "Заказчик", 10, False, 80)
    _center(doc, obj["name"], 16, True)
    _center(doc, obj["address"], 12, False, 30)
    _center(doc, volume_title.upper(), 16, True)
    _center(doc, f"Шифр: {volume_code}", 12, False, 120)
    _center(doc, f"Москва · {date.today().year}", 12)
    doc.add_page_break()


def build_id_volume(db, obj, checklist, plan) -> bytes:
    doc = _doc()
    code = f"{obj['code']}-ИД"
    _title_page(doc, obj, "Исполнительная документация", code)

    doc.add_heading("Опись документов тома", level=1)
    _table(doc, ["№", "Наименование документа", "Примечание"],
           [(i + 1, n, "") for i, n in enumerate(ID_ORDER[1:])])
    doc.add_page_break()

    doc.add_heading("Акт технического обследования", level=1)
    doc.add_paragraph(
        f"Объект: {obj['name']}, {obj['address']}. Тип: {BUILDING_TYPES.get(obj['building_type'])}. "
        f"Год постройки: {obj['year']}. Этажность: {obj['floors']}. Общая площадь: {obj['area']} м².")
    c = checklist.get("constructions", {})
    _table(doc, ["Конструкция", "Техническое состояние"],
           [(CONSTR_RU[k], STATE_RU.get(c.get(k, ""), "—")) for k in CONSTR_RU])
    if c.get("notes"):
        doc.add_paragraph(f"Замечания: {c['notes']}")
    works = checklist.get("works", [])
    if works:
        _table(doc, ["Вид работ", "Отметка"], [(w["name"], "Выполнено" if w["done"] else "Не выполнено") for w in works])
    doc.add_page_break()

    doc.add_heading("Ведомость инженерных систем и узлов", level=1)
    sys_rows = []
    for s in obj["systems"]:
        info = checklist.get("systems", {}).get(s, {})
        sys_rows.append((SYSTEMS[s], STATE_RU.get(info.get("state", ""), "—"), info.get("notes", "")))
    _table(doc, ["Система", "Состояние", "Примечание"], sys_rows)
    nodes = [n for n in db["nodes"] if n["id"] in checklist.get("nodes", [])]
    if nodes:
        _table(doc, ["Система", "Тип", "Узел"], [(SYSTEMS[n["system"]], n["type"], n["name"]) for n in nodes])
    doc.add_page_break()

    doc.add_heading("Обмерные планы", level=1)
    doc.add_paragraph(f"Масштаб {plan['scale']}. Площадь типового этажа {plan['floor_area']} м². "
                      "Чертёж формируется по данным лазерного сканирования.")
    _table(doc, ["№", "Помещение", "Площадь, м²"],
           [(i + 1, r["name"], r["area"]) for i, r in enumerate(plan["rooms"])])
    doc.add_page_break()

    doc.add_heading("Ведомость применённых материалов", level=1)
    mats = [m for m in db["materials"] if m["id"] in checklist.get("materials", [])]
    _table(doc, ["№", "Материал", "Поставщик", "Сертификат / документ", "Действует до"],
           [(i + 1, m["name"], m["supplier"], m["cert_no"], m["cert_until"]) for i, m in enumerate(mats)])

    doc.add_heading("Сертификаты и паспорта", level=1)
    if mats:
        for m in mats:
            doc.add_paragraph(f"[Скан сертификата] {m['name']} — {m['cert_no']}"
                              + (" + паспорт изделия" if m["passport"] else ""), style="List Bullet")
    else:
        doc.add_paragraph("Материалы в чек-листе не выбраны.")

    doc.add_heading("Протоколы лабораторных испытаний", level=1)
    tickets = [t for t in db["lab_tickets"] if t["object_id"] == obj["id"]]
    if tickets:
        for t in tickets:
            doc.add_paragraph(f"{t['test_type']} — {t['lab']}, заявка от {t['created']}, статус: {t['status']}",
                              style="List Bullet")
    else:
        doc.add_paragraph("Лабораторные испытания не требовались.")
    return _save(doc)


RD_SECTIONS = {
    "AR": ("Архитектурные решения", "АР"),
    "KR": ("Конструктивные решения", "КР"),
    "OV": ("Отопление, вентиляция и кондиционирование", "ОВ"),
    "VK": ("Водоснабжение и водоотведение", "ВК"),
    "EOM": ("Электроснабжение и электроосвещение", "ЭОМ"),
    "TS": ("Тепловые сети и ИТП", "ТС"),
}


def build_rd_section(db, obj, checklist, plan, section: str) -> bytes:
    title, mark = RD_SECTIONS[section]
    doc = _doc()
    _title_page(doc, obj, f"Рабочая документация. Раздел «{title}»", f"{obj['code']}-{mark}")
    doc.add_heading("Общие данные", level=1)
    doc.add_paragraph(
        f"Рабочая документация разработана на основании технического обследования и лазерного "
        f"сканирования объекта «{obj['name']}». Оформление — по ГОСТ Р 21.101-2020.")
    _table(doc, ["Лист", "Наименование", "Примечание"], [
        (1, "Общие данные", ""), (2, "Обмерный план типового этажа", "по данным сканирования"),
        (3, "Ведомость существующего оборудования", "из чек-листа"), (4, "Решения (заготовка)", "дополнить проектировщику"),
    ])
    doc.add_heading("Исходное состояние (из чек-листа)", level=1)
    if section in ("AR", "KR"):
        c = checklist.get("constructions", {})
        for k, n in CONSTR_RU.items():
            doc.add_paragraph(f"{n}: {STATE_RU.get(c.get(k, ''), '—')}", style="List Bullet")
        if c.get("notes"):
            doc.add_paragraph(c["notes"])
    else:
        info = checklist.get("systems", {}).get(section, {})
        doc.add_paragraph(f"Состояние системы: {STATE_RU.get(info.get('state', ''), '—')}. {info.get('notes', '')}")
        nodes = [n for n in db["nodes"] if n["id"] in checklist.get("nodes", []) and n["system"] == section]
        for n in nodes:
            doc.add_paragraph(f"{n['type']}: {n['name']}", style="List Bullet")
    doc.add_heading("Проектные решения", level=1)
    doc.add_paragraph("[Заготовка. Раздел дополняется проектировщиком.]")
    return _save(doc)


def build_access_request(obj, items: list[str]) -> bytes:
    doc = _doc()
    p = doc.add_paragraph(f"{COMPANY}\nИсх. № ____ от {date.today():%d.%m.%Y}")
    p.paragraph_format.space_after = Pt(18)
    r = doc.add_paragraph(f"Руководителю\n{obj['customer']}")
    r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _center(doc, "ЗАЯВКА НА ДОПУСК", 14, True, 12)
    doc.add_paragraph(
        f"В рамках выполнения работ по техническому обследованию и лазерному сканированию объекта "
        f"«{obj['name']}» по адресу: {obj['address']}, просим обеспечить доступ специалистов "
        f"{COMPANY} в следующие помещения:")
    for it in items:
        doc.add_paragraph(it, style="List Number")
    doc.add_paragraph("Также просим обеспечить присутствие ответственного представителя для открытия "
                      "электрощитовых и технических помещений.")
    doc.add_paragraph("\n\nГенеральный директор  ____________________ /                 /")
    return _save(doc)


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
